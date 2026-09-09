"""Material controlado para probar los cuatro mundos mediante importación PPTX."""
import json
from pathlib import Path
from pptx import Presentation
from practice_fixtures import SOURCES, TECH_SOURCE

out=Path('/tmp/mindcrafted-adventure-check');out.mkdir(exist_ok=True)
presentation=Presentation()
notes={'technology':TECH_SOURCE,**SOURCES}
titles={'technology':'Redes y algoritmos','mathematics':'Igualdades y fracciones','laboratory':'Corriente y energía','biology':'La vida de una semilla'}
for world,source in notes.items():
    slide=presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text=titles[world];slide.placeholders[1].text=source
presentation.save(out/'expedicion.pptx')
(out/'notes.json').write_text(json.dumps(notes,ensure_ascii=False),encoding='utf-8')
print(out)
