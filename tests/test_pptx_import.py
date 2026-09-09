import io
import unittest

from pptx import Presentation

from mindcrafted.pptx_import import (
    MAX_GAME_SECTIONS,
    PPTXImportError,
    import_pptx,
    slides_to_markdown,
    validate_pptx,
)
from mindcrafted.server import parse_markdown


class PowerPointImportTests(unittest.TestCase):
    def test_real_pptx_is_converted_to_game_ready_markdown(self):
        presentation = Presentation()
        presentation.core_properties.title = "Curso de Astronomía"
        for title, points in (
            ("El sistema solar", ["El Sol es una estrella", "La Tierra es un planeta"]),
            ("Las órbitas", ["La gravedad mantiene los planetas en órbita", "Las órbitas son elípticas"]),
        ):
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = title
            frame = slide.placeholders[1].text_frame
            frame.clear()
            for index, point in enumerate(points):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = point
        buffer = io.BytesIO()
        presentation.save(buffer)

        result = import_pptx(buffer.getvalue(), "astronomia.pptx")
        parsed = parse_markdown(result.markdown)

        self.assertEqual(result.slides, 2)
        self.assertEqual(result.sections, 2)
        self.assertEqual(parsed["course"]["title"], "Curso de Astronomía")
        self.assertEqual(len(parsed["chunks"]), 2)
        self.assertIn("gravedad", parsed["chunks"][1]["content"])

    def test_many_slides_are_grouped_into_bounded_game_sections(self):
        slides = [
            {"number": index, "title": f"Tema {index}", "blocks": ["Explicación educativa suficientemente detallada para la lección."]}
            for index in range(1, 31)
        ]
        markdown = slides_to_markdown(slides, "curso.pptx")

        self.assertLessEqual(markdown.count("\n## "), MAX_GAME_SECTIONS)
        self.assertIn("### Tema 30", markdown)

    def test_rejects_legacy_or_invalid_powerpoint(self):
        with self.assertRaisesRegex(PPTXImportError, "antiguo"):
            validate_pptx(b"data", "curso.ppt")
        with self.assertRaisesRegex(PPTXImportError, "estructura"):
            validate_pptx(b"not-a-zip", "curso.pptx")


if __name__ == "__main__":
    unittest.main()
