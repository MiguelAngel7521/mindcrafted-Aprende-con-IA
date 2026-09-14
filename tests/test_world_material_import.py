"""Keep document → lesson boundaries covered during the runtime migration."""
import io

import pytest
from pptx import Presentation

from mindcrafted.pdf_import import PDFImportError, pages_to_markdown, validate_pdf
from mindcrafted.pptx_import import PPTXImportError, import_pptx, validate_pptx
from mindcrafted.server import _validate_generation_content, parse_markdown


def test_pdf_pages_remain_usable_world_source():
    text = pages_to_markdown([
        "Manual de Redes\nCAPÍTULO 1 - ROUTING\nLos paquetes deben alcanzar su destino mediante una ruta conectada. Cada router tiene una capacidad limitada.\nPágina 1",
        "Manual de Redes\nCAPÍTULO 2 - CONGESTIÓN\nDistribuir la carga entre routers evita la congestión. Los paquetes se pierden si el nodo supera su capacidad.\nPágina 2",
    ], "redes.pdf")
    content = parse_markdown(text)
    content["course"]["gameplay"] = "world"
    assert _validate_generation_content(content)
    assert "Página 1" not in text
    assert "capacidad" in " ".join(c["content"] for c in content["chunks"])


def test_pptx_can_supply_region_content():
    presentation = Presentation()
    presentation.core_properties.title = "Ciudad de los Nodos"
    for title, body in [("Routing", "Los paquetes deben alcanzar su destino. Los enlaces sin salida no entregan información al destino."),
                        ("Congestión", "Cada router procesa una cantidad limitada de paquetes. Repartir la carga evita que todos compitan por el mismo camino.")]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    buffer = io.BytesIO(); presentation.save(buffer)
    result = import_pptx(buffer.getvalue(), "redes.pptx")
    content = parse_markdown(result.markdown)
    content["course"]["gameplay"] = "world"
    assert _validate_generation_content(content)
    assert result.slides == 2
    assert len(content["chunks"]) == 2


def test_invalid_pdf_and_ppt_are_rejected():
    with pytest.raises(PDFImportError):
        validate_pdf(b"not a PDF", "source.pdf")
    with pytest.raises(PPTXImportError):
        validate_pptx(b"not a presentation", "source.ppt")


def test_scanned_pdf_requires_text_extraction():
    with pytest.raises(PDFImportError, match="OCR"):
        pages_to_markdown(["1", "2", "3"], "scan.pdf")
