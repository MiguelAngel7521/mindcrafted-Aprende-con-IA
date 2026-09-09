"""Crea un PPTX de prueba en /tmp para probar la subida real por HTTP."""
from pathlib import Path
import json
from pptx import Presentation
from practice_fixtures import SOURCES
out=Path('/tmp/mindcrafted-practice-check');out.mkdir(exist_ok=True)
pres=Presentation()
for world, source in SOURCES.items():
    slide=pres.slides.add_slide(pres.slide_layouts[1])
    slide.shapes.title.text=world
    slide.placeholders[1].text=source
pres.save(out/'apuntes.pptx')
(out/'sources.json').write_text(json.dumps(SOURCES, ensure_ascii=False))
print(out)
