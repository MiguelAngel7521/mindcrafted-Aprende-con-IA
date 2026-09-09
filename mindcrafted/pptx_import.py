"""Conversión local de presentaciones PowerPoint (.pptx) a Markdown educativo."""

from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass
from pathlib import Path


MAX_PPTX_BYTES = 25 * 1024 * 1024
MAX_PPTX_SLIDES = 200
MAX_MARKDOWN_CHARS = 300_000
MAX_GAME_SECTIONS = 12


class PPTXImportError(ValueError):
    """Error de PowerPoint apto para mostrar al usuario."""


@dataclass(frozen=True)
class PPTXImportResult:
    markdown: str
    filename: str
    slides: int
    characters: int
    sections: int
    truncated: bool = False

    def as_dict(self) -> dict:
        return {
            "markdown": self.markdown,
            "filename": self.filename,
            "slides": self.slides,
            "characters": self.characters,
            "sections": self.sections,
            "truncated": self.truncated,
        }


def _safe_filename(filename: str) -> str:
    basename = Path((filename or "presentacion.pptx").replace("\\", "/")).name
    basename = re.sub(r"[\x00-\x1f]+", " ", basename).strip()
    return basename[:220] or "presentacion.pptx"


def _clean(value: object, limit: int = 2_000) -> str:
    text = re.sub(r"[ \t]+", " ", str(value or "")).strip()
    return text[:limit]


def validate_pptx(data: bytes, filename: str) -> None:
    if not filename.lower().endswith(".pptx"):
        if filename.lower().endswith(".ppt"):
            raise PPTXImportError("El formato .ppt antiguo no es compatible; guárdalo como .pptx y vuelve a subirlo")
        raise PPTXImportError("El archivo debe tener extensión .pptx")
    if not data:
        raise PPTXImportError("La presentación está vacía")
    if len(data) > MAX_PPTX_BYTES:
        raise PPTXImportError(f"La presentación supera el límite de {MAX_PPTX_BYTES // (1024 * 1024)} MB")
    if not data.startswith(b"PK"):
        raise PPTXImportError("El archivo no tiene una estructura PowerPoint .pptx válida")


def _text_frame_lines(text_frame) -> list[str]:
    paragraphs = []
    for paragraph in text_frame.paragraphs:
        text = _clean(paragraph.text)
        if not text:
            continue
        level = min(max(int(getattr(paragraph, "level", 0) or 0), 0), 5)
        paragraphs.append((level, text))
    if not paragraphs:
        return []
    if len(paragraphs) == 1 and paragraphs[0][0] == 0:
        return [paragraphs[0][1]]
    return [("  " * level) + "- " + text for level, text in paragraphs]


def _table_markdown(table) -> str:
    rows = [[_clean(cell.text, 500).replace("|", "\\|") for cell in row.cells] for row in table.rows]
    if not rows or not rows[0]:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = rows[0]
    output = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    output.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n".join(output)


def extract_slides(data: bytes) -> tuple[list[dict], str]:
    try:
        from pptx import Presentation
        from pptx.exc import PackageNotFoundError
    except ImportError as exc:  # pragma: no cover - depende del entorno
        raise PPTXImportError(
            "Falta el lector de PowerPoint. Instala las dependencias con: pip install -r requirements.txt"
        ) from exc

    try:
        presentation = Presentation(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError, OSError) as exc:
        raise PPTXImportError("No se pudo leer la presentación; puede estar dañada") from exc
    if len(presentation.slides) > MAX_PPTX_SLIDES:
        raise PPTXImportError(f"La presentación supera el límite de {MAX_PPTX_SLIDES} diapositivas")

    slides: list[dict] = []
    for number, slide in enumerate(presentation.slides, 1):
        title_shape = slide.shapes.title
        title = _clean(title_shape.text, 180) if title_shape is not None else ""
        blocks: list[str] = []
        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if getattr(shape, "has_table", False):
                table = _table_markdown(shape.table)
                if table:
                    blocks.append(table)
                continue
            if getattr(shape, "has_text_frame", False):
                lines = _text_frame_lines(shape.text_frame)
                if lines:
                    blocks.append("\n".join(lines))
        if title or blocks:
            slides.append({"number": number, "title": title or f"Diapositiva {number}", "blocks": blocks})

    metadata_title = _clean(getattr(presentation.core_properties, "title", ""), 180)
    return slides, metadata_title


def slides_to_markdown(slides: list[dict], filename: str, metadata_title: str = "") -> str:
    filename = _safe_filename(filename)
    if not slides:
        raise PPTXImportError("La presentación no contiene texto legible")
    text_size = sum(len(str(block)) for slide in slides for block in slide.get("blocks", []))
    if text_size < 60:
        raise PPTXImportError("No se encontró texto suficiente en la presentación")

    course_title = _clean(metadata_title, 180) or _clean(slides[0].get("title"), 180) or Path(filename).stem
    group_size = max(1, math.ceil(len(slides) / min(MAX_GAME_SECTIONS, len(slides))))
    output = [f"# {course_title}", "", f"Contenido importado de `{filename}`."]
    for offset in range(0, len(slides), group_size):
        group = slides[offset:offset + group_size]
        section_title = _clean(group[0].get("title"), 180) or f"Diapositivas {offset + 1}-{offset + len(group)}"
        output.extend(["", f"## {section_title}"])
        for position, slide in enumerate(group):
            slide_title = _clean(slide.get("title"), 180)
            if position > 0 and slide_title:
                output.extend(["", f"### {slide_title}"])
            blocks = [_clean(block, 20_000) for block in slide.get("blocks", []) if _clean(block, 20_000)]
            if blocks:
                output.extend(["", "\n\n".join(blocks)])
    return "\n".join(output).strip() + "\n"


def import_pptx(data: bytes, filename: str) -> PPTXImportResult:
    filename = _safe_filename(filename)
    validate_pptx(data, filename)
    slides, metadata_title = extract_slides(data)
    markdown = slides_to_markdown(slides, filename, metadata_title)
    truncated = len(markdown) > MAX_MARKDOWN_CHARS
    if truncated:
        markdown = markdown[:MAX_MARKDOWN_CHARS].rsplit("\n", 1)[0]
        markdown += "\n\n> Contenido truncado por el límite de importación.\n"
    return PPTXImportResult(
        markdown=markdown,
        filename=filename,
        slides=len(slides),
        characters=len(markdown),
        sections=len(re.findall(r"^## ", markdown, re.MULTILINE)),
        truncated=truncated,
    )
